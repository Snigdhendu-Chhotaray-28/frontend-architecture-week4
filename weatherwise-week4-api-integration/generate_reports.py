"""
Report Generator Script: Converts Markdown Report into Word (.docx) and PDF (.pdf) formats.
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.pdfgen import canvas

BASE_DIR = r"d:\Users\DELL\WebDev\Projects\Git projects\yuva\frontend-architecture-week4\weatherwise-week4-api-integration"
MD_PATH = os.path.join(BASE_DIR, "Week_4_API_Integration_and_Data_Handling_Report.md")
DOCX_PATH = os.path.join(BASE_DIR, "Week_4_API_Integration_and_Data_Handling_Report.docx")
PDF_PATH = os.path.join(BASE_DIR, "Week_4_API_Integration_and_Data_Handling_Report.pdf")

def read_markdown_file():
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        return f.read()

def generate_docx():
    print("Generating DOCX report...")
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Title & Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("Week 4 Internship Report: API Integration and Data Handling\n")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(14, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Project: WeatherWise — Modern Real-Time Weather Dashboard\n")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = RGBColor(2, 132, 199)
    sub_run.font.bold = True

    # Cover Meta Box
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Shading Accent 1'
    meta_data = [
        ("Assignment Name", "Week 4: API Integration and Data Handling in React.js"),
        ("Track", "Frontend Architecture & React.js Web Development"),
        ("Student Name", "[Student Name]"),
        ("Institution", "[Institution / Organization Name]"),
        ("Submission Date", "October 9, 2026"),
        ("Tech Stack", "React 18, Vite 6, Vanilla CSS3, Lucide React, Open-Meteo APIs")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = table.rows[i]
        row.cells[0].paragraphs[0].text = k
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[1].paragraphs[0].text = v

    doc.add_page_break()

    # Read markdown lines
    md_content = read_markdown_file()
    lines = md_content.split('\n')

    in_code_block = False
    code_text = []

    for line in lines:
        stripped = line.strip()

        if stripped.startswith('```'):
            if in_code_block:
                # End of code block
                code_p = doc.add_paragraph()
                code_p.paragraph_format.left_indent = Inches(0.2)
                code_run = code_p.add_run('\n'.join(code_text))
                code_run.font.name = 'Consolas'
                code_run.font.size = Pt(9.5)
                code_run.font.color.rgb = RGBColor(30, 41, 59)
                code_text = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_text.append(line)
            continue

        if not stripped:
            continue

        if stripped.startswith('# '):
            continue # Already handled in cover
        elif stripped.startswith('## '):
            p = doc.add_heading(stripped[3:], level=1)
            p.runs[0].font.name = 'Arial'
            p.runs[0].font.color.rgb = RGBColor(14, 23, 42)
        elif stripped.startswith('### '):
            p = doc.add_heading(stripped[4:], level=2)
            p.runs[0].font.name = 'Arial'
            p.runs[0].font.color.rgb = RGBColor(2, 132, 199)
        elif stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(stripped[2:])
        elif stripped.startswith('|'):
            # Table lines (handled gracefully as text)
            p = doc.add_paragraph()
            p.add_run(stripped)
            p.runs[0].font.name = 'Consolas'
            p.runs[0].font.size = Pt(8.5)
        else:
            p = doc.add_paragraph()
            p.add_run(stripped)

    doc.save(DOCX_PATH)
    print(f"DOCX created successfully at: {DOCX_PATH}")

def generate_pdf():
    print("Generating PDF report...")
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0e172a'),
        alignment=1, # Center
        spaceAfter=8
    )

    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#0284c7'),
        alignment=1,
        spaceAfter=15
    )

    h1_style = ParagraphStyle(
        'DocH1',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#0f172a'),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#0369a1'),
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    )

    code_style = ParagraphStyle(
        'DocCode',
        parent=styles['Code'],
        fontName='Courier',
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor('#1e293b'),
        backColor=colors.HexColor('#f1f5f9'),
        spaceBefore=4,
        spaceAfter=6
    )

    story = []

    # Title & Cover Header
    story.append(Paragraph("Week 4 Internship Report: API Integration and Data Handling", title_style))
    story.append(Paragraph("Project: WeatherWise — Modern Real-Time Weather Dashboard", sub_style))
    story.append(Spacer(1, 10))

    # Meta Table
    meta_table_data = [
        [Paragraph("<b>Assignment</b>", body_style), Paragraph("Week 4: API Integration and Data Handling", body_style)],
        [Paragraph("<b>Track</b>", body_style), Paragraph("Frontend Architecture & React.js Web Development", body_style)],
        [Paragraph("<b>Student Name</b>", body_style), Paragraph("[Student Name]", body_style)],
        [Paragraph("<b>Institution</b>", body_style), Paragraph("[Institution / Organization Name]", body_style)],
        [Paragraph("<b>Submission Date</b>", body_style), Paragraph("October 9, 2026", body_style)],
        [Paragraph("<b>Tech Stack</b>", body_style), Paragraph("React 18, Vite 6, Vanilla CSS3, Lucide React, Open-Meteo APIs", body_style)]
    ]
    t = Table(meta_table_data, colWidths=[130, 390])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t)
    story.append(Spacer(1, 15))

    md_content = read_markdown_file()
    lines = md_content.split('\n')

    in_code = False
    code_lines = []

    for line in lines:
        stripped = line.strip()

        if stripped.startswith('```'):
            if in_code:
                code_text = '<br/>'.join([c.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;') for c in code_lines])
                story.append(Paragraph(code_text, code_style))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not stripped:
            continue

        if stripped.startswith('# '):
            continue
        elif stripped.startswith('## '):
            story.append(Paragraph(stripped[3:], h1_style))
        elif stripped.startswith('### '):
            story.append(Paragraph(stripped[4:], h2_style))
        elif stripped.startswith('* ') or stripped.startswith('- '):
            bullet_text = "&bull; " + stripped[2:].replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            story.append(Paragraph(bullet_text, body_style))
        elif stripped.startswith('|'):
            story.append(Paragraph(stripped.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'), code_style))
        else:
            safe_text = stripped.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            # Bold highlights
            safe_text = safe_text.replace('**', '<b>', 1)
            while '**' in safe_text:
                safe_text = safe_text.replace('**', '</b>', 1)
                if '**' in safe_text:
                    safe_text = safe_text.replace('**', '<b>', 1)
            story.append(Paragraph(safe_text, body_style))

    doc.build(story)
    print(f"PDF created successfully at: {PDF_PATH}")

if __name__ == "__main__":
    generate_docx()
    generate_pdf()
